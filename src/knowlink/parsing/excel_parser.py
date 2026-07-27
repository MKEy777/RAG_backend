from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import pandas as pd

from .base import ParseResult, ParserCapabilityError
from .registry import register

DISPIMG_RE = re.compile(r'=_?xlfn\.(DISPIMG|IMAGE)\s*\(', re.IGNORECASE)


def _col_letter(idx: int) -> str:
    result = ""
    while idx >= 0:
        result = chr(65 + idx % 26) + result
        idx = idx // 26 - 1
    return result


def _repair_shared_strings(data: bytes) -> bytes:
    """Repair XLSX packages where sharedStrings.xml is missing or mis-pathed.

    Some WPS/export tools reference xl/sharedStrings.xml in [Content_Types].xml
    or workbook.xml.rels but omit the actual part. If worksheets only use inline
    strings, strip the dangling references so openpyxl can open the file.
    """
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return data

    names = {n.replace("\\", "/") for n in zf.namelist()}

    # sharedStrings.xml exists at the expected path — no repair needed
    if "xl/sharedStrings.xml" in names:
        return data

    # Maybe it exists at a different path (case/path mismatch)
    sst_actual = next((n for n in names if n.lower().endswith("sharedstrings.xml")), None)
    if sst_actual:
        # Rewrite zip with the part renamed to the expected path
        files = {n.replace("\\", "/"): zf.read(n) for n in zf.namelist()}
        files["xl/sharedStrings.xml"] = files.pop(sst_actual)
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for name, content in files.items():
                zout.writestr(name, content)
        return out.getvalue()

    # sharedStrings.xml is truly absent — check if manifest references it
    ct = zf.read("[Content_Types].xml").decode("utf-8", errors="replace") if "[Content_Types].xml" in names else ""
    rels_path = "xl/_rels/workbook.xml.rels"
    rels = zf.read(rels_path).decode("utf-8", errors="replace") if rels_path in names else ""
    if "sharedstrings" not in ct.lower() and "sharedstrings" not in rels.lower():
        return data  # not referenced, nothing to fix

    # Check if any worksheet actually uses shared string cells (t="s")
    uses_sst = False
    for n in names:
        if n.startswith("xl/worksheets/") and n.endswith(".xml"):
            sheet = zf.read(n).decode("utf-8", errors="replace")
            if 't="s"' in sheet:
                uses_sst = True
                break
    if uses_sst:
        return data  # genuinely needs sharedStrings but it's missing — unrepairable

    # Strip dangling references so openpyxl can open
    import re as _re
    files = {n.replace("\\", "/"): zf.read(n) for n in zf.namelist()}
    if "[Content_Types].xml" in files:
        ct_text = files["[Content_Types].xml"].decode("utf-8")
        ct_text = _re.sub(r'<Override[^>]*PartName="[^"]*sharedStrings\.xml"[^>]*/>', "", ct_text, flags=_re.IGNORECASE)
        files["[Content_Types].xml"] = ct_text.encode("utf-8")
    if rels_path in files:
        rels_text = files[rels_path].decode("utf-8")
        rels_text = _re.sub(r'<Relationship[^>]*Type="[^"]*sharedStrings"[^>]*/>', "", rels_text, flags=_re.IGNORECASE)
        files[rels_path] = rels_text.encode("utf-8")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, content in files.items():
            zout.writestr(name, content)
    return out.getvalue()


def _fill_merged_cells_xlsx(data: bytes) -> bytes:
    """Unmerge ranges and copy master cell value into every covered cell.

    openpyxl only stores values on the top-left cell of a merge; pandas then
    sees NaN in the rest. Filling at the openpyxl level is precise — only
    truly merged cells get filled, unlike df.ffill() which blindly propagates.
    """
    try:
        if not zipfile.is_zipfile(io.BytesIO(data)):
            return data
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), data_only=True)
        changed = False
        for ws in wb.worksheets:
            if not ws.merged_cells.ranges:
                continue
            for merge_range in list(ws.merged_cells.ranges):
                master_value = ws.cell(merge_range.min_row, merge_range.min_col).value
                ws.unmerge_cells(str(merge_range))
                for row in range(merge_range.min_row, merge_range.max_row + 1):
                    for col in range(merge_range.min_col, merge_range.max_col + 1):
                        ws.cell(row, col).value = master_value
                changed = True
        if not changed:
            return data
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    except Exception:
        return data


def _filter_dispimg(text: str) -> str:
    return DISPIMG_RE.sub("", text).strip()


_HEADER_END_PUNCT = set("。？！；.?!;")


def _extract_headers(df: pd.DataFrame) -> tuple[list[str], bool]:
    """Extract column names from the first row.

    Returns (headers, has_header_row). If the first row looks like data
    (cells too long or ending with sentence punctuation), treat it as data
    and fall back to column letters.
    """
    if df.empty:
        return [], False
    first_row = df.iloc[0]
    texts: list[str] = []
    for col_idx, val in enumerate(first_row):
        if pd.isna(val):
            texts.append("")
        else:
            texts.append(_filter_dispimg(str(val)).strip())

    non_empty = [t for t in texts if t]
    if not non_empty:
        return [_col_letter(i) for i in range(len(texts))], False

    avg_len = sum(len(t) for t in non_empty) / len(non_empty)
    ends_with_punct = any(t[-1] in _HEADER_END_PUNCT for t in non_empty if t)

    # Heuristic: real headers are short labels, not sentences
    if avg_len > 15 or ends_with_punct:
        return [_col_letter(i) for i in range(len(texts))], False

    headers = [t if t else _col_letter(i) for i, t in enumerate(texts)]
    return headers, True


def _parse_xlsx_data(data: bytes) -> ParseResult:
    data = _repair_shared_strings(data)
    data = _fill_merged_cells_xlsx(data)
    try:
        xls = pd.ExcelFile(io.BytesIO(data), engine="openpyxl")
    except Exception:
        from .docx_parser import _convert_with_libreoffice

        converted = _convert_with_libreoffice(data, ".xlsx", ".xlsx")
        if converted is None:
            raise ParserCapabilityError("xlsx", "LibreOffice (soffice)")
        xls = pd.ExcelFile(io.BytesIO(converted), engine="openpyxl")

    parts: list[str] = []
    for sheet_name in xls.sheet_names:
        df = xls.parse(sheet_name, header=None)
        if df.empty:
            continue
        headers, has_header = _extract_headers(df)
        parts.append(f"## Sheet: {sheet_name}\n")
        for row_idx, row in df.iterrows():
            if has_header and row_idx == 0:
                continue
            cells = []
            all_empty = True
            for col_idx, val in enumerate(row):
                if pd.isna(val):
                    continue
                text = _filter_dispimg(str(val))
                if not text:
                    continue
                all_empty = False
                col_name = headers[col_idx] if col_idx < len(headers) else _col_letter(col_idx)
                cells.append(f"{col_name}: {text}")
            if all_empty:
                continue
            parts.append(", ".join(cells))

    return ParseResult(markdown="\n".join(parts))


@register("xlsx")
def parse_xlsx(data: bytes, filename: str, **kwargs) -> ParseResult:
    return _parse_xlsx_data(data)


@register("xlsb")
def parse_xlsb(data: bytes, filename: str, **kwargs) -> ParseResult:
    try:
        xls = pd.ExcelFile(io.BytesIO(data), engine="pyxlsb")
        parts: list[str] = []
        for sheet_name in xls.sheet_names:
            df = xls.parse(sheet_name, header=None)
            if df.empty:
                continue
            headers, has_header = _extract_headers(df)
            parts.append(f"## Sheet: {sheet_name}\n")
            for row_idx, row in df.iterrows():
                if has_header and row_idx == 0:
                    continue
                cells = []
                all_empty = True
                for col_idx, val in enumerate(row):
                    if pd.isna(val):
                        continue
                    text = str(val).strip()
                    if not text:
                        continue
                    all_empty = False
                    col_name = headers[col_idx] if col_idx < len(headers) else _col_letter(col_idx)
                    cells.append(f"{col_name}: {text}")
                if all_empty:
                    continue
                parts.append(", ".join(cells))
        return ParseResult(markdown="\n".join(parts))
    except Exception:
        from .docx_parser import _convert_with_libreoffice

        converted = _convert_with_libreoffice(data, ".xlsb", ".xlsx")
        if converted is not None:
            return _parse_xlsx_data(converted)
        raise ParserCapabilityError("xlsb", "pyxlsb or LibreOffice")


@register("ods")
def parse_ods(data: bytes, filename: str, **kwargs) -> ParseResult:
    from .docx_parser import _convert_with_libreoffice

    converted = _convert_with_libreoffice(data, ".ods", ".xlsx")
    if converted is not None:
        return _parse_xlsx_data(converted)
    try:
        xls = pd.ExcelFile(io.BytesIO(data), engine="odf")
        parts: list[str] = []
        for sheet_name in xls.sheet_names:
            df = xls.parse(sheet_name, header=None)
            if df.empty:
                continue
            headers, has_header = _extract_headers(df)
            parts.append(f"## Sheet: {sheet_name}\n")
            for row_idx, row in df.iterrows():
                if has_header and row_idx == 0:
                    continue
                cells = []
                all_empty = True
                for col_idx, val in enumerate(row):
                    if pd.isna(val):
                        continue
                    text = str(val).strip()
                    if not text:
                        continue
                    all_empty = False
                    col_name = headers[col_idx] if col_idx < len(headers) else _col_letter(col_idx)
                    cells.append(f"{col_name}: {text}")
                if all_empty:
                    continue
                parts.append(", ".join(cells))
        return ParseResult(markdown="\n".join(parts))
    except Exception:
        raise ParserCapabilityError("ods", "LibreOffice (soffice) or odfpy")


@register("xls")
def parse_xls(data: bytes, filename: str, **kwargs) -> ParseResult:
    from .docx_parser import _convert_with_libreoffice

    converted = _convert_with_libreoffice(data, ".xls", ".xlsx")
    if converted is not None:
        return _parse_xlsx_data(converted)
    try:
        xls = pd.ExcelFile(io.BytesIO(data), engine="xlrd")
        parts: list[str] = []
        for sheet_name in xls.sheet_names:
            df = xls.parse(sheet_name, header=None)
            if df.empty:
                continue
            headers, has_header = _extract_headers(df)
            parts.append(f"## Sheet: {sheet_name}\n")
            for row_idx, row in df.iterrows():
                if has_header and row_idx == 0:
                    continue
                cells = []
                all_empty = True
                for col_idx, val in enumerate(row):
                    if pd.isna(val):
                        continue
                    text = str(val).strip()
                    if not text:
                        continue
                    all_empty = False
                    col_name = headers[col_idx] if col_idx < len(headers) else _col_letter(col_idx)
                    cells.append(f"{col_name}: {text}")
                if all_empty:
                    continue
                parts.append(", ".join(cells))
        return ParseResult(markdown="\n".join(parts))
    except Exception:
        raise ParserCapabilityError("xls", "LibreOffice (soffice)")
