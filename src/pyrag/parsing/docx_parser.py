from __future__ import annotations

import io
import subprocess
import tempfile
from pathlib import Path

from .base import ImageRef, ParseResult, ParserCapabilityError
from .registry import OLE_MAGIC, register


def _has_libreoffice() -> bool:
    try:
        subprocess.run(
            ["soffice", "--version"],
            capture_output=True,
            timeout=10,
        )
        return True
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return False


def _convert_with_libreoffice(data: bytes, suffix: str, target_suffix: str) -> bytes | None:
    if not _has_libreoffice():
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / f"input{suffix}"
        src.write_bytes(data)
        try:
            subprocess.run(
                ["soffice", "--headless", "--convert-to", target_suffix.lstrip("."), "--outdir", tmp, str(src)],
                capture_output=True,
                timeout=120,
            )
        except (subprocess.TimeoutExpired, FileNotFoundError):
            return None
        out = Path(tmp) / f"input{target_suffix}"
        if out.exists():
            return out.read_bytes()
    return None


def _parse_docx_markitdown(data: bytes) -> ParseResult | None:
    try:
        from markitdown import MarkItDown

        md = MarkItDown()
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            f.write(data)
            f.flush()
            result = md.convert(f.name)
        text = result.text_content if result else ""
        if not text or not text.strip():
            return None
        return ParseResult(markdown=text)
    except Exception:
        return None


def _parse_docx_python_docx(data: bytes) -> ParseResult:
    from docx import Document

    doc = Document(io.BytesIO(data))
    images: list[ImageRef] = []
    parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = (para.style.name or "").lower()
        if style_name.startswith("heading"):
            try:
                level = int(style_name.replace("heading", "").strip())
            except ValueError:
                level = 1
            parts.append(f"{'#' * level} {text}")
        else:
            parts.append(text)

    for table in doc.tables:
        rows_text = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows_text.append("| " + " | ".join(cells) + " |")
        if rows_text:
            header_sep = "| " + " | ".join(["---"] * len(table.rows[0].cells)) + " |"
            rows_text.insert(1, header_sep)
            parts.append("\n".join(rows_text))

    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            try:
                img_data = rel.target_part.blob
                images.append(ImageRef(data=img_data, mime="image/png"))
            except Exception:
                pass

    return ParseResult(markdown="\n\n".join(parts), images=images)


@register("docx")
def parse_docx(data: bytes, filename: str) -> ParseResult:
    if data[:8] == OLE_MAGIC:
        return parse_doc(data, filename)

    result = _parse_docx_markitdown(data)
    if result is not None:
        return result
    return _parse_docx_python_docx(data)


@register("doc")
def parse_doc(data: bytes, filename: str) -> ParseResult:
    converted = _convert_with_libreoffice(data, ".doc", ".docx")
    if converted is not None:
        result = _parse_docx_markitdown(converted)
        if result is not None:
            return result
        return _parse_docx_python_docx(converted)

    if not _has_libreoffice():
        raise ParserCapabilityError("doc", "LibreOffice (soffice)")

    raise ParserCapabilityError("doc", "LibreOffice (soffice)")
